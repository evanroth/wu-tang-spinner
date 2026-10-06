"""Export the "Keep" rows of wu-tang-spinners.xlsx to spinner-verbs.json and SOURCES.md.

Run from the project folder after editing the spreadsheet:
    python3 tools/export_verbs.py
Needs openpyxl (pip install openpyxl).
"""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "wu-tang-spinners.xlsx"
OUT = ROOT / "spinner-verbs.json"
SOURCES = ROOT / "SOURCES.md"

ws = load_workbook(SHEET, read_only=True)["Spinners"]
header = [c.value for c in next(ws.iter_rows(max_row=1))]
col = {name: header.index(name) for name in ("Spinner text", "Status", "Song", "Album", "Who")}

verbs, table = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    text = (row[col["Spinner text"]] or "").strip()
    if text and (row[col["Status"]] or "").strip().lower() == "keep" and text not in verbs:
        verbs.append(text)
        cells = [text] + [str(row[col[k]] or "").strip() for k in ("Song", "Album", "Who")]
        table.append("| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |")

OUT.write_text(json.dumps({"mode": "replace", "verbs": verbs}, indent=2, ensure_ascii=False) + "\n")
SOURCES.write_text(
    "# Sources\n\n"
    "Where each spinner word comes from: Wu-Tang Clan song titles, hooks and short lines.\n"
    "Generated from `wu-tang-spinners.xlsx` by `tools/export_verbs.py`.\n\n"
    "| Spinner text | Song | Album | Who |\n|---|---|---|---|\n" + "\n".join(table) + "\n"
)
print(f"Wrote {len(verbs)} verbs to {OUT.name} and {SOURCES.name}")
